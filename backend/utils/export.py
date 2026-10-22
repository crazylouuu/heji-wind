"""
ChemELN Export utilities — PDF/Markdown/JSON/SQLite backup.
"""
import base64
import mimetypes
import json
import os
import re
import sqlite3
from datetime import datetime
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from openpyxl import Workbook
from utils.chem import render_structure_image
from config import DB_PATH, ATTACH_DIR, BACKUP_DIR

# ------------------------------------------------------------------ 统一字体
# 三种导出（HTML / PDF / Word）统一为「西文 Times New Roman + 中文 宋体」。
#   · HTML / Word 由渲染器逐字符回退，给一条字体链即可；
#   · PDF（xhtml2pdf）只能做「元素级」回退（getFontName 取第一个已注册的名字），
#     所以中文片段必须单独包一层 span —— 见 export_formats.wrap_font_runs()。
FONT_LATIN = "Times New Roman"
FONT_CJK = "宋体"
FONT_CJK_EN = "SimSun"          # 宋体的英文名（非中文 Windows 上按此名查找）
FONT_SYMBOL = "Segoe UI Symbol"  # 兜底：⟶ ⇌ ✕ 等 Times 与宋体都缺的符号
FONT_STACK = (f'"{FONT_LATIN}", "{FONT_CJK}", "{FONT_CJK_EN}", '
              f'"{FONT_SYMBOL}", serif')

# -------------------------------------- 投料质量 / 摩尔数 / 体积 显示规则
# 投料表三列数值（投料质量(g)、摩尔(mmol)、体积(mL)）按数值大小决定显示位数：
#   v < 0 → 3 位小数 ｜ 0 ≤ v < 10 → 2 位小数 ｜ v ≥ 10 → 1 位小数
# ⛔ 这是**展示层**规则：数据库与化学计量计算仍用完整精度，
#    与 Excel「单元格值不变、只是数字格式变了」是同一个模型。
#    前端同名实现见 frontend/js/format.js（Fmt.qty），两边必须同步；
#    verify_export.py 里有跨语言一致性断言。
QTY_EXCEL_FORMAT = "[<0]0.000;[>=10]0.0;0.00"   # Excel 自定义格式（最多两个条件）
QTY_EXCEL_COLS = (6, 8)                         # 投料表 sheet 里需套该格式的列（F..H = 质量/摩尔/体积）


def fmt_quantity(value):
    """把投料表数值（质量/摩尔/体积）格式化成显示字符串（None / 非数值原样返回）。"""
    if value is None or value == "":
        return value
    try:
        v = Decimal(str(value))
    except (InvalidOperation, ValueError, TypeError):
        return value
    if not v.is_finite():
        return value
    if v < 0:
        quantum = Decimal("0.001")
    elif v < 10:
        quantum = Decimal("0.01")
    else:
        quantum = Decimal("0.1")
    # 用十进制字面量做半进位（ROUND_HALF_UP），避免二进制浮点的 2.675 → 2.67
    return f"{v.quantize(quantum, rounding=ROUND_HALF_UP):f}"


# style="..." 属性（单双引号都要管）
_STYLE_ATTR = re.compile(r"(style\s*=\s*)([\"'])(.*?)\2", re.I | re.S)
_FONT_FAMILY_DECL = re.compile(r"font-family\s*:[^;]*;?", re.I)
# <font face="...">（老式 execCommand('fontName') 的产物）
_FONT_FACE_ATTR = re.compile(r"(<font\b[^>]*?)\s+face\s*=\s*([\"']).*?\2", re.I)


def strip_inline_fonts(html: str) -> str:
    """去掉正文里的内联字体声明，保证导出字体完全统一。

    ⛔ 只清 font-family / <font face>，字号、颜色、加粗一律保留。
    正文是富文本，理论上编辑器可能插进 `style="font-family:..."`：
    内联样式的优先级高于导出的样式表，会直接把统一字体顶掉。
    """
    if not html:
        return html

    def _fix_style(m):
        css = _FONT_FAMILY_DECL.sub("", m.group(3))
        css = re.sub(r";\s*;+", ";", css).strip().strip(";")
        if not css:
            return ""                      # 整个 style 只剩字体声明 → 连属性一起删
        return f"{m.group(1)}{m.group(2)}{css}{m.group(2)}"

    html = _STYLE_ATTR.sub(_fix_style, html)
    return _FONT_FACE_ATTR.sub(r"\1", html)


def attachment_abs_path(rel_path: str) -> str:
    """把 attachments 下的相对路径（可能含反斜杠）转成绝对路径，不存在则返回 ''。"""
    if not rel_path:
        return ""
    rel = str(rel_path).replace("\\", "/").lstrip("/")
    full = os.path.join(ATTACH_DIR, *rel.split("/"))
    return full if os.path.isfile(full) else ""


def attachment_data_uri(rel_path: str) -> str:
    """把附件图片读成 data URI。

    ⛔ 单文件导出（HTML/PDF）绝不能写 `src="/attachments/xxx"`：导出文件是被
    下载到本地再打开的，该路径会落在文件系统根目录上 → 图片全部裂开。
    必须把图片内嵌进导出文件本身。
    """
    full = attachment_abs_path(rel_path)
    if not full:
        return ""
    mime = mimetypes.guess_type(full)[0] or "image/png"
    try:
        with open(full, "rb") as f:
            return f"data:{mime};base64," + base64.b64encode(f.read()).decode("ascii")
    except OSError:
        return ""


def inline_attachment_images(html: str) -> str:
    """把正文 HTML 里 src="/attachments/..." 的图片替换为内嵌 data URI（找不到文件则原样保留）。"""
    if not html:
        return html

    def _repl(m):
        uri = attachment_data_uri(m.group(1))
        return f'src="{uri}"' if uri else m.group(0)

    return re.sub(r'src="/attachments/([^"]+)"', _repl, html)


def export_experiment_markdown(exp_id: int) -> str:
    """Export a single experiment to Markdown."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    exp = conn.execute("SELECT * FROM experiments WHERE id=?", (exp_id,)).fetchone()
    if not exp:
        conn.close()
        return "# Experiment not found"

    tags = [r["name"] for r in conn.execute(
        "SELECT t.name FROM tags t JOIN experiment_tags et ON t.id=et.tag_id WHERE et.experiment_id=?",
        (exp_id,))]
    reagents = conn.execute("SELECT * FROM reagents WHERE experiment_id=? ORDER BY sort_order",
                            (exp_id,)).fetchall()
    attachments = conn.execute("SELECT * FROM attachments WHERE experiment_id=?",
                               (exp_id,)).fetchall()
    compound = None
    if exp["fk_compound_id"]:
        compound = conn.execute("SELECT * FROM compounds WHERE id=?",
                                (exp["fk_compound_id"],)).fetchone()

    conn.close()

    lines = [
        f"# {exp['title']}",
        "",
        f"**项目:** {exp['project']}  ",
        f"**日期:** {exp['exp_date']}  ",
        f"**状态:** {exp['status']}  ",
        f"**目标化合物:** {exp['target_compound']}  ",
    ]
    if exp["yield_pct"] is not None:
        lines.append(f"**收率:** {exp['yield_pct']}%  ")
    if exp["purity_pct"] is not None:
        lines.append(f"**纯度:** {exp['purity_pct']}%  ")
    if tags:
        lines.append(f"**标签:** {', '.join(tags)}  ")
    lines.append("")
    lines.append("---")
    lines.append("")
    lines.append(exp["content"] or "(无内容)")
    lines.append("")

    if compound:
        lines.append("## 化合物信息")
        lines.append(f"- **名称:** {compound['name']}")
        lines.append(f"- **CAS:** {compound['cas']}")
        lines.append(f"- **SMILES:** `{compound['smiles']}`")
        lines.append(f"- **分子式:** {compound['formula']}")
        lines.append(f"- **分子量:** {compound['mw']}")
        lines.append("")

    if reagents:
        lines.append("## 投料表")
        lines.append("| 试剂 | 角色 | 投料比 | 投料质量(g) | 摩尔(mmol) | 体积(mL) | 纯度 |")
        lines.append("|------|------|--------|-------------|-----------|---------|------|")
        for r in reagents:
            if r["role"] == "substrate":
                ratio = "基准"
            else:
                rv = r["ratio_value"] if r["ratio_value"] is not None else r["equivalents"]
                ratio = f"{rv} wt" if r["ratio_type"] == "weight" else (f"{rv} eq" if rv is not None else "-")
            mass = r["mass_g"]
            if mass is None and r["mass_mg"] is not None:
                mass = r["mass_mg"] / 1000.0
            mol = None if r["ratio_type"] == "weight" else r["moles_mmol"]
            vol = None if r["ratio_type"] == "weight" else r["volume_ml"]
            lines.append(
                f"| {r['name'] or '-'} | {r['role'] or '-'} | "
                f"{ratio} | {fmt_quantity(mass) if mass is not None else '-'} | "
                f"{fmt_quantity(mol) if mol is not None else '-'} | "
                f"{fmt_quantity(vol) if vol is not None else '-'} | "
                f"{r['purity_pct'] or '-'} |")
        lines.append("")

    if attachments:
        lines.append("## 附件")
        for a in attachments:
            lines.append(f"- {a['filename']} ({a['file_type']}) — {a['description']}")
        lines.append("")

    lines.append("---")
    lines.append(f"*导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}*")

    return "\n".join(lines)


def export_experiment_html(exp_id: int) -> str:
    """Export a single experiment to a standalone HTML file.
    The experiment `content` is already rich-text HTML (may contain embedded
    base64 reaction structures and reagent tables), so it is injected directly."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    exp = conn.execute("SELECT * FROM experiments WHERE id=?", (exp_id,)).fetchone()
    if not exp:
        conn.close()
        return "<html><body><h1>Experiment not found</h1></body></html>"

    tags = [r["name"] for r in conn.execute(
        "SELECT t.name FROM tags t JOIN experiment_tags et ON t.id=et.tag_id WHERE et.experiment_id=?",
        (exp_id,))]
    reagents = conn.execute("SELECT * FROM reagents WHERE experiment_id=? ORDER BY sort_order",
                            (exp_id,)).fetchall()
    attachments = conn.execute("SELECT * FROM attachments WHERE experiment_id=?",
                               (exp_id,)).fetchall()
    compound = None
    if exp["fk_compound_id"]:
        compound = conn.execute("SELECT * FROM compounds WHERE id=?",
                                (exp["fk_compound_id"],)).fetchone()
    reaction = conn.execute(
        "SELECT * FROM reactions WHERE experiment_id=? ORDER BY id DESC LIMIT 1",
        (exp_id,)).fetchone()
    conn.close()

    meta_rows = [
        ("项目", exp["project"]), ("日期", exp["exp_date"]), ("状态", exp["status"]),
        ("目标化合物", exp["target_compound"]),
    ]
    if exp["yield_pct"] is not None:
        meta_rows.append(("收率", f'{exp["yield_pct"]}%'))
    if exp["purity_pct"] is not None:
        meta_rows.append(("纯度", f'{exp["purity_pct"]}%'))
    if tags:
        meta_rows.append(("标签", ", ".join(tags)))

    meta_html = "".join(
        f'<tr><th style="text-align:left;padding:4px 10px;color:#64748b;width:120px">{k}</th>'
        f'<td style="padding:4px 10px">{v or "—"}</td></tr>'
        for k, v in meta_rows
    )

    reagent_html = ""
    if reagents:
        role_names = {"substrate": "原料", "reagent": "试剂", "solvent": "溶剂", "product": "产物"}
        def _ratio_cell(r):
            if r["role"] == "substrate":
                return "基准"
            rv = r["ratio_value"] if r["ratio_value"] is not None else r["equivalents"]
            if rv is None:
                return "—"
            return f'{rv} wt' if r["ratio_type"] == "weight" else f'{rv} eq'
        def _mass_cell(r):
            mass = r["mass_g"]
            if mass is None and r["mass_mg"] is not None:
                mass = r["mass_mg"] / 1000.0
            return "—" if mass is None else fmt_quantity(mass)
        def _mol_cell(r):
            if r["ratio_type"] == "weight":
                return "—"
            return "—" if r["moles_mmol"] is None else fmt_quantity(r["moles_mmol"])
        def _vol_cell(r):
            if r["ratio_type"] == "weight":
                return "—"
            return "—" if r["volume_ml"] is None else fmt_quantity(r["volume_ml"])
        rows = "".join(
            f'<tr><td>{i+1}</td><td>{r["name"] or "—"}</td>'
            f'<td>{role_names.get(r["role"], r["role"] or "—")}</td>'
            f'<td>{r["mw"] if r["mw"] is not None else "—"}</td>'
            f'<td>{_ratio_cell(r)}</td>'
            f'<td>{_mass_cell(r)}</td>'
            f'<td>{_mol_cell(r)}</td>'
            f'<td>{_vol_cell(r)}</td></tr>'
            for i, r in enumerate(reagents)
        )
        reagent_html = (
            '<h2>投料表</h2>'
            '<table border="1" cellpadding="6" cellspacing="0" style="border-collapse:collapse;width:100%">'
            '<thead><tr><th>#</th><th>名称</th><th>角色</th><th>MW</th><th>投料比</th>'
            '<th>投料质量(g)</th><th>摩尔(mmol)</th><th>体积(mL)</th></tr></thead>'
            f'<tbody>{rows}</tbody></table>'
        )

    compound_html = ""
    if compound:
        compound_html = (
            '<h2>化合物信息</h2>'
            f'<ul><li><strong>名称:</strong> {compound["name"]}</li>'
            f'<li><strong>CAS:</strong> {compound["cas"] or "—"}</li>'
            f'<li><strong>SMILES:</strong> <code>{compound["smiles"] or ""}</code></li>'
            f'<li><strong>分子式:</strong> {compound["formula"] or "—"}</li>'
            f'<li><strong>分子量:</strong> {compound["mw"] or "—"}</li></ul>'
        )

    # Reaction equation (rendered from the reaction row fetched above)
    reaction_html = ""
    if reaction:
        if (reaction["display_mode"] or "scheme") == "image" and reaction["image_path"]:
            # 图片模式：必须内嵌，导出文件脱离服务后 /attachments/ 路径会失效
            uri = attachment_data_uri(reaction["image_path"])
            if uri:
                reaction_html = ('<h2>反应方程式</h2>'
                                 f'<div class="rxn-image-wrap"><img src="{uri}" alt="reaction"></div>')
            else:
                reaction_html = ('<h2>反应方程式</h2>'
                                 '<p style="color:#b91c1c">反应式图片文件缺失，无法显示</p>')
        else:
            try:
                reactants = json.loads(reaction["reactant_smiles"] or "[]")
                products = json.loads(reaction["product_smiles"] or "[]")
            except (json.JSONDecodeError, TypeError):
                reactants, products = [], []
            if reactants or products:
                def mol_block(it):
                    img = ""
                    if it.get("smiles"):
                        img = render_structure_image(it["smiles"], size=200)
                    name = it.get("name") or ""
                    meta = ""
                    if it.get("mw"):
                        meta = f'MW {it["mw"]}' + (f' · {it["formula"]}' if it.get("formula") else "")
                    if img:
                        return (f'<div class="rxn-mol"><img src="data:image/png;base64,{img}" alt="">'
                                f'<div class="rxn-mol-name">{name}</div>'
                                f'<div class="rxn-mol-meta">{meta}</div></div>')
                    if name:
                        return f'<div class="rxn-mol rxn-mol-err"><div class="rxn-mol-name">{name}</div><div class="rxn-mol-meta">未渲染</div></div>'
                    return f'<div class="rxn-mol rxn-mol-err">❌ {it.get("smiles","")}</div>'

                def side(items):
                    return "".join(mol_block(it) for it in items)

                arrow = reaction["arrow_label"] or ""
                cond = reaction["conditions"] or ""
                reaction_html = (
                    '<h2>反应方程式</h2><div class="rxn-equation">'
                    f'<div class="rxn-side">{side(reactants)}</div>'
                    '<div class="rxn-arrow-wrap">'
                    f'<div class="rxn-arrow-label">{arrow}</div>'
                    '<div class="rxn-arrow">⟶</div>'
                    f'<div class="rxn-cond">⚙ {cond}</div></div>'
                    f'<div class="rxn-side">{side(products)}</div></div>'
                )

    # 正文里插过的图片/PDF 图标同样要内嵌，否则下载后是裂图
    content_html = inline_attachment_images(
        strip_inline_fonts(exp["content"] or "<p>（无内容）</p>"))

    return f"""<!DOCTYPE html>
<html lang="zh-CN"><head><meta charset="utf-8">
<title>{exp['title']}</title>
<style>
  html, body, div, p, span, li, dt, dd, td, th, h1, h2, h3, h4, h5, h6,
  b, i, u, em, strong, code, pre, blockquote, a, small, sub, sup, ul, ol,
  dl, table, caption, label, font {{ font-family: {FONT_STACK}; }}
  body {{ max-width: 900px; margin: 40px auto; padding: 20px; color:#1e293b; line-height:1.6; }}
  h1,h2 {{ color:#2c3e50; border-bottom:1px solid #e2e8f0; padding-bottom:6px; }}
  table {{ border-collapse: collapse; width: 100%; margin: 12px 0; }}
  th,td {{ border: 1px solid #cbd5e1; padding: 6px 10px; }}
  code {{ background:#f1f5f9; padding:2px 5px; border-radius:3px; }}
  .rxn-equation {{ display:flex; align-items:center; gap:16px; flex-wrap:wrap; }}
  .rxn-side {{ display:flex; align-items:center; gap:10px; flex-wrap:wrap; }}
  .rxn-mol {{ text-align:center; }}
  .rxn-mol img {{ max-height:120px; }}
  .rxn-mol-name {{ font-size:12px; margin-top:4px; }}
  .rxn-mol-meta {{ font-size:11px; color:#64748b; }}
  .rxn-image-wrap {{ text-align:center; }}
  .rxn-image-wrap img {{ max-width:100%; }}
  .rxn-arrow {{ font-size:28px; color:#475569; }}
  .rxn-arrow-label {{ font-size:12px; text-align:center; color:#475569; }}
  .rxn-cond {{ font-size:12px; text-align:center; color:#64748b; }}
</style></head><body>
<h1>{exp['title']}</h1>
<table style="background:#f8fafc">{meta_html}</table>
{reaction_html}
{reagent_html}
<h2>实验内容</h2>
{content_html}
{compound_html}
<hr>
<p style="color:#94a3b8;font-size:12px">导出时间: {datetime.now().strftime('%Y-%m-%d %H:%M')}</p>
</body></html>"""


def export_all_json() -> str:
    """Export the entire database to a JSON string."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    data = {}
    for table in ["experiments", "compounds", "reagents", "templates",
                  "tags", "experiment_tags", "attachments", "versions",
                  "experiment_links", "settings"]:
        try:
            rows = conn.execute(f"SELECT * FROM {table}").fetchall()
            data[table] = [dict(r) for r in rows]
        except Exception:
            data[table] = []

    conn.close()
    data["_export_time"] = datetime.now().isoformat()
    data["_version"] = "1.0"
    return json.dumps(data, ensure_ascii=False, indent=2)


def export_experiments_excel(exp_ids: list = None) -> str:
    """Export experiments to an Excel file. Returns file path."""
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    if exp_ids:
        placeholders = ",".join("?" * len(exp_ids))
        experiments = conn.execute(
            f"SELECT * FROM experiments WHERE id IN ({placeholders})", exp_ids).fetchall()
    else:
        experiments = conn.execute("SELECT * FROM experiments ORDER BY id").fetchall()

    wb = Workbook()
    ws = wb.active
    ws.title = "实验列表"
    headers = ["ID", "标题", "项目", "状态", "日期", "目标化合物",
               "收率(%)", "纯度(%)", "创建时间", "更新时间"]
    ws.append(headers)

    for exp in experiments:
        ws.append([exp["id"], exp["title"], exp["project"], exp["status"],
                   exp["exp_date"], exp["target_compound"],
                   exp["yield_pct"], exp["purity_pct"],
                   exp["created_at"], exp["updated_at"]])

    # Auto-width
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        ws.column_dimensions[col[0].column_letter].width = min(max_len + 2, 40)

    # Second sheet: reagents
    ws2 = wb.create_sheet("投料表")
    ws2.append(["实验ID", "实验标题", "试剂名称", "角色", "投料比",
                "投料质量(g)", "摩尔(mmol)", "体积(mL)", "纯度(%)"])
    for exp in experiments:
        reagents = conn.execute(
            "SELECT * FROM reagents WHERE experiment_id=? ORDER BY sort_order",
            (exp["id"],)).fetchall()
        for r in reagents:
            if r["role"] == "substrate":
                ratio = "基准"
            else:
                rv = r["ratio_value"] if r["ratio_value"] is not None else r["equivalents"]
                ratio = f"{rv} wt" if r["ratio_type"] == "weight" else (f"{rv} eq" if rv is not None else None)
            mass = r["mass_g"]
            if mass is None and r["mass_mg"] is not None:
                mass = round(r["mass_mg"] / 1000.0, 4)
            mol = None if r["ratio_type"] == "weight" else r["moles_mmol"]
            vol = None if r["ratio_type"] == "weight" else r["volume_ml"]
            ws2.append([exp["id"], exp["title"], r["name"], r["role"], ratio,
                        mass, mol, vol, r["purity_pct"]])

    # 投料质量(g) / 摩尔(mmol) / 体积(mL)：保留数值本身（还能继续求和、再计算），
    # 只把单元格数字格式设成「<0 → 3 位；0~10 → 2 位；≥10 → 1 位」，
    # 与 MD / HTML / PDF / Word 的文本格式规则一致。
    for row in ws2.iter_rows(min_row=2, min_col=QTY_EXCEL_COLS[0], max_col=QTY_EXCEL_COLS[1]):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = QTY_EXCEL_FORMAT

    conn.close()

    os.makedirs(BACKUP_DIR, exist_ok=True)
    filepath = os.path.join(BACKUP_DIR, f"experiments_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
    wb.save(filepath)
    return filepath


def backup_database() -> str:
    """Create a SQLite backup copy. Returns file path."""
    os.makedirs(BACKUP_DIR, exist_ok=True)
    backup_name = f"chemeln_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.db"
    backup_path = os.path.join(BACKUP_DIR, backup_name)

    src = sqlite3.connect(DB_PATH)
    dst = sqlite3.connect(backup_path)
    src.backup(dst)
    src.close()
    dst.close()
    return backup_path


def import_from_csv(csv_path: str, table: str = "experiments") -> dict:
    """Import data from CSV into a table. Returns count of imported rows."""
    import csv
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row

    with open(csv_path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        rows = list(reader)

    if not rows:
        conn.close()
        return {"imported": 0, "error": "No data rows in CSV"}

    # Map CSV columns to table columns
    columns = [k for k in rows[0].keys() if k]
    placeholders = ",".join(["?"] * len(columns))
    col_str = ",".join(columns)

    count = 0
    for row in rows:
        values = [row[c] for c in columns]
        try:
            conn.execute(f"INSERT INTO {table} ({col_str}) VALUES ({placeholders})", values)
            count += 1
        except Exception as e:
            # Skip invalid rows
            pass

    conn.commit()
    conn.close()
    return {"imported": count, "table": table}
